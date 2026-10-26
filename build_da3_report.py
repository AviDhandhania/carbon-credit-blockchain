"""Build docs/DA3_Report.docx from docs/DA3_Report.md.

The markdown is the single source of truth; this script only renders it. Body
paragraphs are justified (the DA3 format asks for justified alignment), code
blocks stay monospaced and left-aligned, and tables and figures are carried over.

Run: python build_da3_report.py
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(ROOT, "docs", "DA3_Report.md")
OUT = os.path.join(ROOT, "docs", "DA3_Report.docx")
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x4B, 0x55, 0x63)


def add_inline(paragraph, text):
    """Add **bold**, `code` and *italic* runs to one paragraph."""
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)
    return paragraph


def justify(paragraph, space_after=6):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.15
    return paragraph


def code_block(doc, lines):
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        # Light grey shading so the block reads as code, not prose.
        shd = p._p.get_or_add_pPr()
        elem = shd.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): "F3F4F6"})
        shd.append(elem)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows):
    if len(rows) < 2:
        return
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, cell in enumerate(header):
        paragraph = table.rows[0].cells[i].paragraphs[0]
        add_inline(paragraph, cell)
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt(9)
    for row in body:
        cells = table.add_row().cells
        for i, value in enumerate(row[: len(header)]):
            paragraph = cells[i].paragraphs[0]
            add_inline(paragraph, value.replace("`", ""))
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def cells_of(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build():
    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.9)
        section.left_margin = section.right_margin = Inches(0.95)

    # ---- cover -------------------------------------------------------------
    i = 0
    if lines and lines[0].startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lines[0][2:])
        r.bold = True
        r.font.size = Pt(22)
        r.font.color.rgb = NAVY
        i = 1
    if i < len(lines) and lines[i].startswith("## "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lines[i][3:])
        r.bold = True
        r.font.size = Pt(15)
        r.font.color.rgb = NAVY
        i += 1
    while i < len(lines) and lines[i].strip() != "---":
        if lines[i].strip():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, lines[i].strip())
            for run in p.runs:
                run.font.size = Pt(11)
                run.font.color.rgb = GREY
        i += 1
    i += 1  # skip the first horizontal rule
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    in_code = False
    code_lines = []
    table_rows = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            flush_table()
            i += 1
            continue

        if is_table_row(stripped):
            cells = cells_of(stripped)
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                i += 1
                continue
            table_rows.append(cells)
            i += 1
            continue

        flush_table()

        if stripped == "---":
            i += 1
            continue

        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            alt, path = image.group(1), image.group(2)
            # Markdown image paths are relative to the source document (docs/).
            image_path = os.path.normpath(
                os.path.join(os.path.dirname(SRC), path.replace("/", os.sep))
            )
            if os.path.exists(image_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(image_path, width=Inches(6.4))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(alt)
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = GREY
            i += 1
            continue

        if stripped.startswith("#"):
            md_level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[md_level:].strip()
            # The cover consumes the markdown # and ## levels, so a report
            # section (## 1. Introduction) becomes Word Heading 1.
            h = doc.add_heading(level=max(1, min(md_level - 1, 3)))
            add_inline(h, text)
            for run in h.runs:
                run.font.color.rgb = NAVY
                run.font.name = "Calibri"
            i += 1
            continue

        # Caption line such as *Table 1. ...*
        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(stripped.strip("*"))
            r.italic = True
            r.font.size = Pt(9)
            r.font.color.rgb = GREY
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            justify(p, space_after=3)
        elif numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
            justify(p, space_after=3)
        elif re.match(r"^\[\d+\]", stripped):
            p = doc.add_paragraph()
            add_inline(p, stripped)
            justify(p, space_after=4)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.4)
        else:
            p = doc.add_paragraph()
            add_inline(p, stripped)
            justify(p)

        i += 1

    if in_code and code_lines:
        code_block(doc, code_lines)

    # footer page numbers
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run()
        fld = run._r.makeelement(qn("w:fldSimple"), {qn("w:instr"): "PAGE"})
        run._r.addnext(fld)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
