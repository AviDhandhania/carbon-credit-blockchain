"""Build DA1_Report.docx and DA1_Presentation.pptx from fixed content + diagram PNGs.
Run: python build_deliverables.py
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.text import PP_ALIGN

ROOT = os.path.dirname(os.path.abspath(__file__))
DIAG = os.path.join(ROOT, "diagrams")
NAVY = RGBColor(0x1f, 0x3a, 0x5f)
TEAL = RGBColor(0x2a, 0x9d, 0x8f)
PNAVY, PTEAL, PWHITE, PDARK = PRGB(0x1f, 0x3a, 0x5f), PRGB(0x2a, 0x9d, 0x8f), PRGB(0xff, 0xff, 0xff), PRGB(0x22, 0x22, 0x22)

def img(name):  # diagram path
    return os.path.join(DIAG, name)

# =====================================================================  DOCX
def build_docx():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)

    def h(text, level, color=NAVY):
        p = d.add_heading(text, level)
        for r in p.runs: r.font.color.rgb = color
        return p

    def para(text, bold=False, italic=False, size=11):
        p = d.add_paragraph()
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        return p

    def lead(label, text):
        """Bold lead-in followed by normal text, in one paragraph."""
        p = d.add_paragraph()
        r = p.add_run(label); r.bold = True; r.font.size = Pt(11)
        r2 = p.add_run(text); r2.font.size = Pt(11)
        return p

    def bullets(items):
        for it in items:
            p = d.add_paragraph(style="List Bullet"); p.add_run(it)

    def table(headers, rows):
        t = d.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
        for i, hd in enumerate(headers):
            c = t.rows[0].cells[i]; c.text = ""
            r = c.paragraphs[0].add_run(hd); r.bold = True; r.font.size = Pt(10)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(val)); r.font.size = Pt(9.5)
        return t

    # Title page
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Blockchain Based Carbon Credit Trading System"); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY
    s = d.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = s.add_run("Digital Assignment 1 (DA1), Blockchain Technology"); rs.font.size = Pt(14); rs.italic = True
    sg = d.add_paragraph(); sg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rg = sg.add_run("SDG 13, Climate Action  |  also supports SDG 7 and SDG 12"); rg.font.size = Pt(12); rg.font.color.rgb = TEAL; rg.bold = True
    d.add_paragraph()
    au = d.add_paragraph(); au.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ra = au.add_run("Submitted by"); ra.font.size = Pt(12); ra.italic = True
    a1 = d.add_paragraph(); a1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = a1.add_run("Avi Dhandhania, 25BCE1207"); r1.font.size = Pt(14); r1.bold = True; r1.font.color.rgb = NAVY
    a2 = d.add_paragraph(); a2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = a2.add_run("Shivesh Kumar, 25BCE1067"); r2.font.size = Pt(14); r2.bold = True; r2.font.color.rgb = NAVY
    d.add_paragraph()

    # ------------------------------------------------------------------ 1
    h("1. Problem Statement, Objectives and Scope", 1)
    h("1.1 The problem", 2)
    para("One carbon credit is permission to release one tonne of CO2. A company that pollutes less than its "
         "allowance can sell the credits it did not use. A company that pollutes more has to buy them. On paper "
         "this pays companies to cut emissions, and it should work.")
    para("In practice the market has a trust problem. The same four issues keep coming up.")
    lead("The same credit gets counted twice. ",
         "Registries are separate databases run by separate bodies, and they do not talk to each other. A credit "
         "sold in one registry can be sold again in another, and nobody notices.")
    lead("Buyers cannot check what they bought. ",
         "If a company buys 10,000 credits from a tree planting project, it has no way to confirm that the trees "
         "exist. It has to trust a PDF.")
    lead("Some credits are for projects that never happened. ",
         "These are called phantom credits. Journalists have found cases where a large share of a registry's "
         "credits did not represent any real reduction. Because the records sit in private databases, this only "
         "comes out years later, if at all.")
    lead("Checking is slow and expensive. ",
         "Approval and settlement still run on manual audits and paperwork. It can take months, and the fees come "
         "out of money that was supposed to go into climate projects.")
    para("The result is that a system built to reduce emissions ends up letting companies claim reductions that "
         "never happened.")

    h("1.2 Objectives", 2)
    bullets([
        "O1. Turn each carbon credit into a token on a public blockchain so it has one identity and cannot be copied.",
        "O2. Handle issuing, selling and retiring credits in smart contract code, so double counting is blocked by "
        "the rules of the system and not by someone remembering to check.",
        "O3. Keep the whole history of every credit on a ledger that anyone can read and nobody can edit.",
        "O4. Cut out middlemen so verification is faster and cheaper.",
        "O5. Give the four kinds of users (project developer, verifier, buyer, regulator) separate on chain roles, "
        "so only the right person can do each action.",
        "O6. Make verifiers pay for bad approvals. A verifier has to lock money before signing off on a project, "
        "and loses it if the project is later shown to be fake. This is the part we have not seen in existing "
        "work, and it is described in Section 4.4.",
    ])

    h("1.3 Scope", 2)
    lead("What DA1 covers: ",
         "the design of a decentralised app (DApp) on Ethereum. This includes an ERC-20 token for credits, a "
         "marketplace contract for listing and buying, a retirement contract that burns credits and issues a "
         "certificate, a staking and challenge contract for verifiers, role based access control, IPFS for project "
         "documents, and a dashboard where anyone can look up a credit.")
    lead("What DA1 does not cover: ",
         "connecting to government registries, legally binding settlement, buying credits with real money, and any "
         "machine learning estimate of how much CO2 a project actually saved. DA1 is the problem study, the "
         "literature review, the justification for using blockchain, the architecture and the plan. Writing the "
         "contracts comes later.")

    h("1.4 How this maps to the SDGs", 2)
    table(["SDG", "What the project does for it"], [
        ["SDG 13, Climate Action (main)", "Makes offsets trustworthy, so money spent on credits funds real emission cuts."],
        ["SDG 7, Affordable and Clean Energy", "Solar and wind projects can sell credits that buyers actually believe in."],
        ["SDG 12, Responsible Consumption", "Public records make greenwashing easy to spot."],
    ])

    # ------------------------------------------------------------------ 2
    h("2. Literature Survey and Research Gap", 1)
    h("2.1 What already exists", 2)
    table(["Work", "What it did well", "Where it falls short"], [
        ["Toucan Protocol", "Moved real credits onto Ethereum as tokens and added on chain retirement.",
         "Everything depends on trusting the bridge. It was criticised for bringing old, low quality credits on "
         "chain, which put good tokens and junk tokens in the same pool."],
        ["KlimaDAO", "Used tokenised credits as a treasury asset to push the carbon price up.",
         "The goal is financial, not honest accounting. It does not check whether a credit is real."],
        ["IBM and Energy Web pilots", "Enterprise blockchains for tracking company emissions.",
         "Permissioned, so only members can read the ledger. The public and regulators are locked out, and that is "
         "the group that most needs to check."],
        ["Verra and Gold Standard", "Well developed methods for measuring a project, plus human auditors.",
         "Central databases that do not reconcile with each other. Double counting and slow settlement are still normal."],
        ["Academic surveys, 2019 to 2023", "Showed that smart contract trading is workable and modelled parts of it.",
         "Mostly stay at concept level. Few cover the full mint, trade and retire path with role control and "
         "evidence anchoring together."],
        ["EU ETS and economic models", "Model how an emissions trading scheme behaves as a market.",
         "Assume the records are correct. They do not deal with the trust layer at all."],
    ])

    h("2.2 The gaps we found", 2)
    lead("G1. The lifecycle is handled in pieces. ",
         "Most systems tokenise credits, or trade them, or retire them. Few enforce the order, so the rule that a "
         "retired credit can never be sold again is often a policy rather than something the code refuses to do.")
    lead("G2. Off chain evidence is loosely attached. ",
         "When a credit is bridged from an old registry it carries the same weaknesses it had there. The verifier's "
         "signature and the project documents are usually kept somewhere else, so a token on its own tells you "
         "very little about where it came from.")
    lead("G3. Public chains have transparency but weak governance. ",
         "Public systems let anyone read the ledger, but they rarely combine that with strict on chain roles for "
         "developer, verifier, buyer and regulator.")
    lead("G4. Nothing happens to a verifier who approves a fake project. ",
         "This gap is the important one. Every design above assumes the data entering the chain is honest. If a "
         "verifier signs off on a project that does not exist, the blockchain records that lie perfectly and "
         "forever. The verifier faces no on chain consequence, and buyers who paid for those credits have no way "
         "to get anything back. Papers name this the garbage in, garbage out problem and then move on.")
    para("What we do about them. We design one public Ethereum system that keeps the lifecycle order in contract "
         "code (G1), stores the verifier's signature and an IPFS hash of the evidence with every batch of credits "
         "(G2), and puts role based access control on top of a public ledger (G3). For G4 we add a stake and "
         "challenge mechanism, plus batch tracing and a non transferable retirement certificate. Section 4.4 "
         "explains these.", bold=True)

    # ------------------------------------------------------------------ 3
    h("3. Is Blockchain the Right Choice Here?", 1)
    h("3.1 Does this need a blockchain at all?", 2)
    para("It is worth asking, because plenty of projects use a blockchain where a database would do. The usual test "
         "is four questions. Do several parties who do not trust each other need to share the same data? Should no "
         "single organisation own that data? Is a record that cannot be edited necessary? Does removing middlemen help?")
    para("Carbon trading answers yes to all four. Polluters, project developers, verifiers, regulators and the "
         "public all have different interests and do not fully trust one another. The record has to survive "
         "attempts to edit it, because the whole point is to prove what happened. And cutting out slow "
         "intermediaries is one of our objectives.")
    para("There is also a simpler argument. A single registry with one owner and one database is exactly what the "
         "market has now, and that is the thing that failed.")

    h("3.2 Ethereum or Hyperledger Fabric?", 2)
    table(["Criterion", "Ethereum (public)", "Hyperledger Fabric (permissioned)"], [
        ["Who can join", "Anyone can read and verify", "Only invited members"],
        ["Public audit", "Full, which is what we need", "Only within the consortium"],
        ["Tokens", "Built in through ERC-20 and ERC-721", "Needs custom chaincode"],
        ["Trust model", "Trustless", "Members trust each other"],
        ["Cost", "Gas fees, reducible with Layer 2", "No gas, but servers to run"],
        ["Speed", "Slower on L1, fast on L2", "Fast"],
    ])

    h("3.3 We chose public Ethereum", 2)
    para("The deciding factor is who needs to check the records. Credits have to be verifiable by regulators, "
         "journalists, NGOs and ordinary people, not just by the companies in a consortium. That rules out a "
         "permissioned ledger straight away.")
    para("Three other reasons back it up. ERC-20 and ERC-721 already match what we need, where one credit is one "
         "token that cannot be duplicated. The tooling around Ethereum is the most mature, so Solidity, "
         "OpenZeppelin, Hardhat, MetaMask and free testnets are all available and well tested, which matters for "
         "security. And Layer 2 networks like Polygon and Arbitrum bring gas costs down far enough that the cost "
         "argument stops being a real objection.")
    para("Hyperledger would be the better pick if credits were traded quietly among a fixed set of known "
         "companies. That is the opposite of what we want.", italic=True)

    # ------------------------------------------------------------------ 4
    h("4. Architecture, Workflow and Design", 1)
    h("4.1 System architecture", 2)
    para("The system has four layers.")
    bullets([
        "Presentation layer. A React web app that talks to the chain through Web3.js, MetaMask for signing "
        "transactions, and a dashboard for verifiers and regulators.",
        "Application layer. A Node.js REST API, a role manager that maps a wallet address to a role, and an "
        "adapter that pulls in off chain data such as satellite or sensor readings for a project.",
        "Blockchain layer on Ethereum. Four contracts. CreditToken is the ERC-20 token and also keeps a record for "
        "each batch. Marketplace handles listing and buying. RetireAndCertify burns credits and issues the "
        "certificate. VerifierStake holds verifier deposits and runs the challenge process.",
        "Storage layer. The chain itself for transaction history, IPFS for large project documents where only the "
        "hash goes on chain, and an ordinary database that caches metadata so the dashboard loads quickly. The "
        "cache holds nothing that matters, so if it is lost it can be rebuilt from the chain.",
    ])
    d.add_picture(img("architecture.png"), width=Inches(6.3))

    h("4.2 The life of a credit", 2)
    bullets([
        "A developer registers a green project, such as reforestation or rooftop solar, and uploads the supporting "
        "documents to IPFS.",
        "A verifier locks a deposit and approves the project. The approval, the verifier's address and the IPFS "
        "hash of the evidence all go on chain.",
        "Credits are minted as a batch. One token equals one tonne of CO2. The batch keeps a link back to the "
        "project and the verifier who approved it.",
        "A company buys credits through the marketplace contract.",
        "To claim the offset, the company retires the credits, which burns them.",
        "The retirement contract issues a certificate to the buyer's address. The certificate cannot be "
        "transferred or sold.",
        "All of this stays readable by anyone. During the challenge window the batch can still be disputed, and "
        "after that it is settled.",
    ])
    para("Because retirement is a burn, a retired credit cannot be sold again. That is not a rule someone has to "
         "enforce, there is simply nothing left to sell.")
    d.add_picture(img("workflow.png"), width=Inches(6.3))

    h("4.3 Sequence for buying and retiring", 2)
    para("The buyer connects a wallet and picks a listing. The DApp calls buyCredit(batchId, qty). The marketplace "
         "contract moves the tokens and the transaction is recorded on Ethereum. Once it is confirmed, the buyer "
         "calls retire(batchId, qty). The tokens are burned and the contract issues the certificate in the same "
         "transaction, so the burn and the proof of the burn cannot come apart. The buyer ends up with a "
         "transaction hash and a certificate that is anchored to their address.")
    d.add_picture(img("sequence.png"), width=Inches(6.0))

    h("4.4 What is new in our design", 2)
    para("Everything in Sections 4.1 to 4.3 is engineering that others have done in some form. This section is the "
         "part we have not found in the systems we surveyed.")
    lead("Verifiers put money at risk. ",
         "Right now a verifier signs a report and walks away. In our design a verifier must lock a deposit in "
         "VerifierStake before approving anything, and the deposit stays locked through a challenge window of "
         "about 90 days per batch. Anyone can challenge a batch during that window by posting a smaller bond and "
         "an IPFS link to their evidence, such as satellite images showing bare land where a forest was claimed. "
         "An address holding the regulator role rules on the challenge. If the challenge succeeds, the verifier's "
         "stake is slashed. Part goes to the challenger, which pays people to look for fraud, and part goes into a "
         "compensation pool for the buyers of that batch. If the challenge fails, the challenger loses their bond, "
         "which stops people filing junk challenges. A verifier whose stake falls below the minimum loses the "
         "ability to approve anything until they top it up.")
    para("This changes what the blockchain is doing. In the systems we reviewed the chain is a very good filing "
         "cabinet, in that it records whatever it is told, including lies. Here the chain also holds the "
         "incentive. Approving a fake project stops being free.")
    d.add_picture(img("novelty.png"), width=Inches(6.3))
    lead("Every credit can be traced back and flagged. ",
         "Each mint creates a batch that stores the project ID, the verifier who approved it and the IPFS hash of "
         "the evidence, so a token is never anonymous. If a project is later proven fake, the contract can mark "
         "every batch that came from it, and the dashboard shows a warning on each of those credits and on every "
         "wallet still holding them. Compare this with a token pool where good and bad credits are mixed together "
         "and one bad project quietly damages the value of the whole pool. Being able to name the affected credits "
         "is what makes compensation possible at all.")
    lead("The retirement certificate cannot be traded. ",
         "When credits are burned, the contract mints an ERC-721 token to the buyer with transfer disabled, which "
         "is a soulbound token. It records how many tonnes were retired, which batch they came from, the date and "
         "who claimed the offset. This matters for two reasons. A company can point an auditor or a customer at a "
         "public certificate instead of a spreadsheet it wrote itself. And because the certificate cannot move, "
         "the same retirement cannot be resold as proof to a second company. Today a burn is only a hole in the "
         "supply and nothing says who the hole belongs to, so two companies can both point at it. Here the claim "
         "has an owner.")
    lead("Together these three close G4. ",
         "Evidence is bound to the credit, someone loses money if the evidence was false, anyone can start that "
         "process, and the offset claim at the end is public and cannot be reused. The chain stops only proving "
         "that a record exists and starts giving people a reason to file honest records in the first place.")
    lead("What we are not claiming. ",
         "Staking does not make fraud impossible. A verifier who profits more than the stake is worth may still "
         "take the risk, so the minimum stake has to scale with the size of the batch, and choosing that number "
         "properly needs work we have not done. Slashing also depends on a regulator role judging challenges, "
         "which is one trusted point in an otherwise trustless design. We think that is the right trade for now, "
         "because a fully automated ruling would need an oracle that can decide whether a forest exists, and that "
         "does not exist yet. A panel of regulator addresses with majority voting is the obvious next step.")

    h("4.5 Other design decisions", 2)
    bullets([
        "ERC-20 for the credits themselves, because credits of the same batch are interchangeable and buyers want "
        "to buy 500 tonnes rather than 500 individual items. The batch record gives us the project identity "
        "without making every tonne a separate NFT.",
        "ERC-721 with transfers disabled for the certificate, since a certificate is genuinely unique and should "
        "never move.",
        "OpenZeppelin AccessControl for the four roles. It is audited and widely used, and role management is "
        "exactly the kind of code we should not write ourselves.",
        "IPFS for documents. A project report can run to hundreds of pages, which is far too expensive to store on "
        "chain. IPFS addresses content by its hash, so storing the hash on chain is enough to prove that a "
        "document has not been changed.",
        "Retirement is an irreversible burn, enforced in the contract, so double spending a retired credit is not "
        "possible.",
    ])

    # ------------------------------------------------------------------ 5
    h("5. Project Planning", 1)
    h("5.1 Timeline and milestones", 2)
    table(["Phase", "Weeks", "Deliverable"], [
        ["Requirement study and literature survey", "W0 to W2", "Problem statement and survey (DA1)"],
        ["System design and architecture", "W1 to W3", "Architecture and design diagrams (DA1)"],
        ["Smart contract development in Solidity", "W3 to W6", "CreditToken and Marketplace deployed"],
        ["Frontend DApp with Web3 integration", "W4 to W7", "Working DApp connected to MetaMask"],
        ["Staking, challenge and certificate module", "W5 to W8", "VerifierStake and RetireAndCertify deployed"],
        ["Oracle and IPFS integration", "W6 to W8", "Evidence hashes anchored on chain"],
        ["Testing on the Sepolia testnet", "W7 to W9", "Full lifecycle demo, including a slashing case"],
        ["Security audit and gas optimisation", "W9 to W10", "Audit notes and optimised contracts"],
        ["Documentation and final demo", "W9 to W11", "Final report and presentation"],
    ])
    d.add_picture(img("gantt.png"), width=Inches(6.3))

    h("5.2 Feasibility", 2)
    lead("Technical: high. ",
         "Every piece already exists and is open source. Solidity with OpenZeppelin, Hardhat, React with Web3.js, "
         "MetaMask, the Sepolia testnet and IPFS. Nothing here is unproven. The staking contract is the most "
         "involved part, but escrow with slashing is a well known pattern in staking and prediction market "
         "contracts, so we have working examples to follow.")
    lead("Economic: high. ",
         "Development costs nothing beyond time, since testnets and all the tooling are free. For a production "
         "deployment the gas cost is the real number, and moving to Polygon or Arbitrum brings it down to cents "
         "per transaction.")
    lead("Operational: medium. ",
         "This is the weak spot. The design needs accredited verifiers who are willing to put up a deposit, and a "
         "regulator who will rule on challenges. Both are people problems, not code problems. Our answer is to "
         "position the system as a transparency layer that sits alongside existing registries rather than a "
         "replacement, so it can start small with a few verifiers who want to prove their credits are better than "
         "average.")
    lead("Schedule: high. ",
         "Eleven weeks fits one semester, and the phases overlap in a way that leaves room if the contract work "
         "runs long.")

    h("5.3 Risks", 2)
    table(["Risk", "What we do about it"], [
        ["Off chain data is falsified before minting",
         "This is the risk the staking design targets. Require more than one independent verifier attestation, "
         "anchor evidence by IPFS hash, and slash the stake if a challenge succeeds."],
        ["Verifier profits more than the stake is worth",
         "Scale the minimum stake with batch size, and keep the deposit locked through the challenge window rather "
         "than releasing it at approval."],
        ["Nobody bothers to file challenges",
         "Pay the challenger part of the slashed stake, so looking for fraud is worth someone's time."],
        ["High gas fees on L1", "Deploy on Layer 2 and batch operations where possible."],
        ["Bugs in the contracts",
         "Use audited OpenZeppelin libraries, run Slither for static analysis, and test on Sepolia before anything "
         "else. The money held in the staking contract makes this the highest priority audit target."],
        ["Regulators are slow to accept it",
         "Present it as an interoperable transparency layer over existing registries, not a competitor."],
    ])

    h("References (indicative)", 1)
    for ref in [
        "Toucan Protocol, documentation on bridging carbon credits on chain.",
        "KlimaDAO whitepaper, tokenised carbon as a reserve asset.",
        "IBM and Energy Web Foundation, enterprise carbon tracking pilots.",
        "Verra and Gold Standard, carbon credit verification methodologies.",
        "Survey papers on blockchain for carbon emission trading, IEEE and Elsevier, 2019 to 2023.",
        "OpenZeppelin, ERC-20, ERC-721 and AccessControl contract libraries.",
        "EIP-5192, minimal soulbound (non transferable) token standard.",
    ]:
        p = d.add_paragraph(style="List Number"); p.add_run(ref)

    out = os.path.join(ROOT, "docs", "DA1_Report.docx")
    d.save(out); return out

# =====================================================================  PPTX
def build_pptx():
    prs = Presentation(); prs.slide_width = PInches(13.333); prs.slide_height = PInches(7.5)
    BLANK = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height

    def bg(slide, color):
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color

    def textbox(slide, l, t, w, h, text, size, color=PDARK, bold=False, align=PP_ALIGN.LEFT, italic=False):
        tb = slide.shapes.add_textbox(PInches(l), PInches(t), PInches(w), PInches(h))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = align
        r = p.add_run(); r.text = text; f = r.font
        f.size = PPt(size); f.bold = bold; f.italic = italic; f.color.rgb = color; f.name = "Calibri"
        return tb

    def bullets_slide(title, items, subtitle=None, size=18):
        s = prs.slides.add_slide(BLANK); bg(s, PWHITE)
        band = s.shapes.add_shape(1, 0, 0, W, PInches(1.15)); band.fill.solid()
        band.fill.fore_color.rgb = PNAVY; band.line.fill.background()
        textbox(s, 0.5, 0.22, 12.3, 0.8, title, 28, PWHITE, bold=True)
        if subtitle:
            textbox(s, 0.5, 1.25, 12.3, 0.5, subtitle, 15, PTEAL, italic=True)
        top = 2.0 if subtitle else 1.5
        tb = s.shapes.add_textbox(PInches(0.7), PInches(top), PInches(12), PInches(5.2))
        tf = tb.text_frame; tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run(); r.text = "•  " + it
            r.font.size = PPt(size); r.font.color.rgb = PDARK; r.font.name = "Calibri"
            p.space_after = PPt(10)
        return s

    def image_slide(title, image, subtitle=None):
        s = prs.slides.add_slide(BLANK); bg(s, PWHITE)
        band = s.shapes.add_shape(1, 0, 0, W, PInches(1.0)); band.fill.solid()
        band.fill.fore_color.rgb = PNAVY; band.line.fill.background()
        textbox(s, 0.5, 0.18, 12.3, 0.7, title, 26, PWHITE, bold=True)
        from PIL import Image
        iw, ih = Image.open(image).size
        maxw, maxh = 11.5, 5.6
        ratio = min(maxw / (iw / 150), maxh / (ih / 150))
        dw, dh = (iw / 150) * ratio, (ih / 150) * ratio
        s.shapes.add_picture(image, PInches((13.333 - dw) / 2), PInches(1.2), PInches(dw), PInches(dh))
        return s

    # Title slide
    s = prs.slides.add_slide(BLANK); bg(s, PNAVY)
    textbox(s, 0.8, 2.2, 11.7, 1.6, "Blockchain Based Carbon Credit Trading System", 40, PWHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, 0.8, 3.9, 11.7, 0.7, "Digital Assignment 1 (DA1)  |  Blockchain Technology", 20, PRGB(0xcf,0xe3,0xdf), align=PP_ALIGN.CENTER)
    textbox(s, 0.8, 4.7, 11.7, 0.7, "SDG 13, Climate Action   |   also supports SDG 7 and SDG 12", 18, PTEAL, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, 0.8, 5.9, 11.7, 0.5, "Avi Dhandhania (25BCE1207)    Shivesh Kumar (25BCE1067)", 17, PWHITE, bold=True, align=PP_ALIGN.CENTER)

    bullets_slide("Agenda", [
        "The problem, our objectives and scope",
        "What already exists, and the gaps we found",
        "Why a blockchain, and why Ethereum",
        "Architecture, workflow and design",
        "What is new in our design: verifiers put money at risk",
        "Plan, feasibility and risks",
    ])

    # 1
    bullets_slide("1. The Problem", [
        "One credit = permission to release one tonne of CO2. Companies that pollute less sell what they did not use.",
        "The same credit gets counted twice, because registries are separate databases that never reconcile.",
        "A buyer cannot check what it bought. It has to trust a PDF.",
        "Phantom credits: credits issued for projects that never happened, found out years later if at all.",
        "Checking is slow and expensive, and the fees eat money meant for climate projects.",
    ], subtitle="A market built to cut emissions is letting companies claim cuts that never happened")

    bullets_slide("1. Objectives and Scope", [
        "O1. One credit becomes one token that cannot be copied.",
        "O2. Mint, sell and retire handled in contract code, so double counting is blocked by the rules.",
        "O3. A full history anyone can read and nobody can edit.",
        "O4. Fewer middlemen, so verification is faster and cheaper.",
        "O5. Separate on chain roles for developer, verifier, buyer and regulator.",
        "O6. Verifiers lock money before approving, and lose it if the project turns out to be fake.",
        "Scope of DA1: design only. No registry integration, no real money, no ML impact estimates.",
    ], subtitle="SDG 13 Climate Action, also supporting SDG 7 and SDG 12", size=16)

    # 2
    bullets_slide("2. What Already Exists", [
        "Toucan Protocol: moved real credits on chain, but everything rests on trusting the bridge.",
        "KlimaDAO: uses credits as a treasury asset. The goal is financial, not honest accounting.",
        "IBM and Energy Web: permissioned tracking, so the public cannot read the ledger.",
        "Verra and Gold Standard: good measurement methods, central databases that do not reconcile.",
        "Academic surveys 2019 to 2023: mostly concept level, rarely the full lifecycle in one design.",
    ])

    bullets_slide("2. The Gaps We Found", [
        "G1. The lifecycle is handled in pieces, so \"cannot sell after retirement\" is a policy, not code.",
        "G2. Evidence is loosely attached. A bridged token tells you little about where it came from.",
        "G3. Public chains give transparency but rarely strict on chain roles.",
        "G4. Nothing happens to a verifier who approves a fake project. The chain records the lie perfectly, "
        "forever, and buyers get nothing back.",
        "G4 is the one we go after. Everything else assumes the data entering the chain is honest.",
    ], subtitle="What the literature leaves unsolved")

    # 3
    bullets_slide("3. Why Blockchain, and Why Ethereum", [
        "The test: parties who do not trust each other, no single owner, records that cannot be edited, middlemen "
        "worth removing. Carbon trading is yes to all four.",
        "One registry with one owner and one database is what the market has now, and that is what failed.",
        "Regulators, journalists and the public need to check the records, which rules out a permissioned ledger.",
        "ERC-20 and ERC-721 already match \"one credit = one token that cannot be duplicated\".",
        "Mature tooling (Solidity, OpenZeppelin, Hardhat, MetaMask, free testnets) means safer, faster building.",
        "Layer 2 (Polygon, Arbitrum) drops gas to cents, so cost is no longer a real objection.",
    ], subtitle="Public Ethereum, not permissioned Hyperledger Fabric", size=16)

    # 4
    image_slide("4. System Architecture", img("architecture.png"))
    image_slide("4. The Life of a Credit", img("workflow.png"))
    image_slide("4. Sequence: Buy and Retire", img("sequence.png"))

    # 4.4 novelty
    image_slide("4. What is New: Verifiers Stake Money", img("novelty.png"))
    bullets_slide("4. What is New, in Three Parts", [
        "Verifiers put money at risk. A deposit is locked before approval and stays locked for a 90 day challenge "
        "window. Anyone can challenge with a bond and evidence. If the challenge wins, the stake is slashed: part "
        "to the challenger, part to a pool for the buyers.",
        "Every credit can be traced and flagged. A batch stores its project, its verifier and the evidence hash, so "
        "a fake project's credits can be named and marked instead of quietly poisoning a shared pool.",
        "The retirement certificate cannot be traded. Burning credits mints a soulbound ERC-721 to the buyer, so an "
        "offset claim has a public owner and cannot be resold as proof to a second company.",
        "The point: the chain stops being only a filing cabinet that records whatever it is told, including lies. "
        "Approving a fake project stops being free.",
    ], subtitle="Closing G4: the gap every other design leaves open", size=15)

    # 5
    image_slide("5. Project Plan", img("gantt.png"))
    bullets_slide("5. Feasibility and Risks", [
        "Technical: high. Every piece is open source and proven. Slashing escrow is a known pattern we can copy.",
        "Economic: high. Testnets and tooling are free, and Layer 2 keeps production gas at cents per transaction.",
        "Operational: medium. Needs verifiers willing to post a deposit and a regulator to rule on challenges. "
        "People problems, not code problems.",
        "Schedule: high. Eleven weeks fits one semester, with overlap for slippage.",
        "Honest limits: a verifier may still profit more than the stake, and a regulator role judging challenges is "
        "one trusted point in a trustless design.",
    ], size=17)

    # Closing
    s = prs.slides.add_slide(BLANK); bg(s, PNAVY)
    textbox(s, 0.8, 2.6, 11.7, 1.2, "Thank You", 44, PWHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, 0.8, 4.0, 11.7, 0.9, "Make the record honest, and make a false record expensive.", 18, PTEAL, italic=True, align=PP_ALIGN.CENTER)

    out = os.path.join(ROOT, "presentation", "DA1_Presentation.pptx")
    prs.save(out); return out

if __name__ == "__main__":
    print("DOCX:", build_docx())
    print("PPTX:", build_pptx())
